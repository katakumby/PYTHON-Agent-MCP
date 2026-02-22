import logging
import os
import sys
from typing import Generator, Tuple, Dict, Any

import docx  # pip install python-docx
import openpyxl  # pip install openpyxl
# Biblioteki zewnętrzne
import pypdf  # pip install pypdf

from buissnes_agent.MetadataModels import FileMetadata
from buissnes_agent.config_loader import settings

logging.basicConfig(level=logging.INFO, stream=sys.stderr)
logger = logging.getLogger(__name__)

class DataLoaderLocalFileLoader:
    """
    Adapter dla plików lokalnych.
    Obsługuje: TXT, MD, XML, JSON, PDF, DOCX, XLSX.
    """

    def __init__(self, directory: str):
        self.directory = os.path.abspath(directory)

    def list_objects(self) -> Generator[str, None, None]:

        allowed_exts = settings.get("chunking.allowed_extensions", [])
        ext_tuple = tuple(allowed_exts)

        if not os.path.exists(self.directory):
            logger.error(f"Katalog nie istnieje: {self.directory}")
            return

        for root, _, files in os.walk(self.directory):
            for file in files:
                ext = os.path.splitext(file)[1].lower()
                if ext in ext_tuple:
                    yield os.path.join(root, file)

    def load_file_with_metadata(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        filename = os.path.basename(file_path)
        ext = os.path.splitext(file_path)[1].lower()

        # 1. Tworzenie obiektu metadanych (Type Safe)
        meta_obj = FileMetadata(
            source=f"file://{file_path}",
            title=filename,
            extension=ext,
            url=f"file://{file_path}",
            domain="local",
            tags=["local", "filesystem"],
            page_number=None  # Cały plik, więc brak konkretnej strony
        )

        content = ""

        try:
            # --- XLSX (Excel) ---
            if ext == ".xlsx":
                try:
                    # data_only=True pobiera wartości, a nie formuły (np. =SUMA(...))
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    text_parts = []
                    for sheet in wb.worksheets:
                        text_parts.append(f"--- Sheet: {sheet.title} ---")
                        for row in sheet.iter_rows(values_only=True):
                            # Łączymy komórki w wierszu spacją, pomijając puste (None)
                            row_text = " ".join([str(cell) for cell in row if cell is not None])
                            if row_text.strip():
                                text_parts.append(row_text)
                    content = "\n".join(text_parts)
                except Exception as e:
                    logger.error(f"Błąd parsowania XLSX {filename}: {e}")
                    return "", {}

            # --- PDF ---
            elif ext == ".pdf":
                try:
                    reader = pypdf.PdfReader(file_path)
                    content = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
                except Exception as e:
                    logger.error(f"Błąd parsowania PDF {filename}: {e}")
                    return "", {}

            # --- DOCX ---
            elif ext == ".docx":
                try:
                    doc = docx.Document(file_path)
                    content = "\n".join([para.text for para in doc.paragraphs])
                except Exception as e:
                    logger.error(f"Błąd parsowania DOCX {filename}: {e}")
                    return "", {}

            # --- TEKSTOWE ---
            else:
                try:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        content = f.read()
                except Exception as e:
                    logger.error(f"Błąd odczytu tekstu {filename}: {e}")
                    return "", {}

            return content, meta_obj.to_dict()

        except Exception as e:
            logger.error(f"Krytyczny błąd przy pliku {file_path}: {e}")
            return "", {}