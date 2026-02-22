import io
import logging
import os
import sys
from typing import Dict, Any, Generator, Tuple

import docx
import openpyxl
# Biblioteki zewnętrzne
import pypdf

from DataLoaderS3Service import DataLoaderS3Service
from buissnes_agent.MetadataModels import FileMetadata

logging.basicConfig(level=logging.INFO, stream=sys.stderr)
logger = logging.getLogger(__name__)

class DataLoaderS3FileLoader:
    def __init__(self, bucket_name: str, prefix: str):
        self.bucket_name = bucket_name
        self.prefix = prefix
        self.s3_service = DataLoaderS3Service()
        logger.info(f"S3FileLoader initialized. Bucket: {bucket_name}")

    def list_objects(self) -> Generator[str, None, None]:
        return self.s3_service.list_objects(self.bucket_name, self.prefix)

    def load_file_with_metadata(self, s3_key: str) -> Tuple[str, Dict[str, Any]]:

        # 1. Logika wyciągania domeny z hierarchii folderów
        key_without_prefix = s3_key
        if self.prefix and s3_key.startswith(self.prefix):
            # Usuwamy prefix konfiguracyjny, żeby znaleźć "logiczną" domenę (podfolder)
            key_without_prefix = s3_key[len(self.prefix):].lstrip("/")

        # Jeśli plik jest w podkatalogu, to nazwa tego katalogu to domena
        domain_name = key_without_prefix.split('/')[0] if '/' in key_without_prefix else "general"

        filename = os.path.basename(s3_key)
        ext = os.path.splitext(s3_key)[1].lower()

        # 2. Tworzenie obiektu metadanych (Type Safe)
        meta_obj = FileMetadata(
            source=f"s3://{self.bucket_name}/{s3_key}",
            title=filename,
            extension=ext,
            url=f"https://{self.bucket_name}.s3.amazonaws.com/{s3_key}",
            domain=domain_name,
            # tags domyślne, page_number None
        )
        content = ""

        try:
            # --- BINARNE (PDF, DOCX, XLSX) ---
            if ext in ['.pdf', '.docx', '.xlsx']:
                try:
                    file_bytes = self.s3_service.download_bytes(self.bucket_name, s3_key)

                    with io.BytesIO(file_bytes) as f:

                        # XLSX
                        if ext == '.xlsx':
                            wb = openpyxl.load_workbook(f, data_only=True)
                            text_parts = []
                            for sheet in wb.worksheets:
                                text_parts.append(f"--- Sheet: {sheet.title} ---")
                                for row in sheet.iter_rows(values_only=True):
                                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                                    if row_text.strip():
                                        text_parts.append(row_text)
                            content = "\n".join(text_parts)

                        # PDF
                        elif ext == '.pdf':
                            reader = pypdf.PdfReader(f)
                            content = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])

                        # DOCX
                        elif ext == '.docx':
                            doc = docx.Document(f)
                            content = "\n".join([para.text for para in doc.paragraphs])

                except Exception as bin_err:
                    logger.error(f"Błąd parsowania pliku binarnego {s3_key}: {bin_err}")
                    return "", {}

            # --- TEKSTOWE ---
            else:
                try:
                    content = self.s3_service.download_text(self.bucket_name, s3_key)
                except Exception as txt_err:
                    logger.error(f"Błąd pobierania tekstu {s3_key}: {txt_err}")
                    return "", {}

            return content, meta_obj.to_dict()

        except Exception as e:
            logger.error(f"Krytyczny błąd przy pliku {s3_key}: {e}")
            return "", {}